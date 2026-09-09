# -*- coding: utf-8 -*-
import os, sys, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from tailor import T
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

RES = json.load(open('data/resultado.json'))
MAN = {m['id']: m for m in json.load(open('data/manifest.json'))}

wb = openpyxl.Workbook()
ws = wb.active; ws.title = "Ofertas"

HDR = ["#","Empresa","Puesto","Ámbito","Ubicación","Modalidad","Publicada","Idioma","Fuente",
       "Salario mín (€)","Salario máx (€)","Salario medio (€)","Origen salario",
       "Coincidencia CV original (%)","Coincidencia CV adaptado (%)","Mejora (pp)","Mejora (%)",
       "Puntos fuertes","Huecos principales","CV adaptado",
       "Enlace de aplicación","Base de la estimación salarial","Aviso"]

hfill = PatternFill("solid", fgColor="1F3864")
hfont = Font(name="Arial", bold=True, color="FFFFFF", size=10)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin,right=thin,top=thin,bottom=thin)

for c,h in enumerate(HDR,1):
    cell = ws.cell(1,c,h); cell.fill=hfill; cell.font=hfont
    cell.alignment=Alignment(vertical="center", wrap_text=True); cell.border=border
ws.row_dimensions[1].height = 34

for i,r in enumerate(RES, start=2):
    m = MAN[r['id']]
    row = [i-1, r['empresa'], r['puesto'], r['ambito'], r['ubicacion'], r['modalidad'], r['publicada'],
           "Español" if r['idioma']=='es' else "Inglés", r['fuente'],
           r['sal_min'], r['sal_max'], None, "Publicado" if r['sal_origen']=='publicado' else "Estimado",
           r['score_orig'], r['score_adap'], None, None,
           " · ".join(r['fuertes'][:3]), " · ".join(r['huecos'][:3]) or "—",
           m['cv'],
           r['url'], r['sal_base'], r.get('alerta',"—")]
    for c,v in enumerate(row,1):
        cell = ws.cell(i,c,v); cell.border=border
        cell.font = Font(name="Arial", size=10)
        cell.alignment = Alignment(vertical="top", wrap_text=(c in (2,3,5,18,19,22,23)))
    # fórmulas (nunca resultados calculados en Python)
    ws.cell(i,12, f"=ROUND(AVERAGE(J{i}:K{i}),0)").font=Font(name="Arial",size=10)
    ws.cell(i,16, f"=ROUND(O{i}-N{i},1)").font=Font(name="Arial",size=10)
    ws.cell(i,17, f"=IFERROR(ROUND((O{i}-N{i})/N{i}*100,1),\"\")").font=Font(name="Arial",size=10)
    for c in (10,11,12): ws.cell(i,c).number_format = '#,##0 "€"'
    for c in (14,15,16,17): ws.cell(i,c).number_format = '0.0'
    ws.cell(i,21).hyperlink = r['url']; ws.cell(i,21).value = "Aplicar en "+r['fuente']
    ws.cell(i,21).font = Font(name="Arial", size=10, color="0563C1", underline="single")
    if r['sal_origen']=='publicado':
        for c in (10,11,13): ws.cell(i,c).fill = PatternFill("solid", fgColor="E2EFDA")
    else:
        for c in (10,11,13): ws.cell(i,c).fill = PatternFill("solid", fgColor="FFF2CC")
    if r.get('alerta'):
        ws.cell(i,23).fill = PatternFill("solid", fgColor="FCE4E4")

n = len(RES)+1
widths = [4,24,38,20,26,22,11,9,13,13,13,14,13,15,15,10,10,34,34,44,20,60,50]
for c,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:W{n}"

# Hoja de método y leyenda
ws2 = wb.create_sheet("Método y leyenda")
lineas = [
 ("CÓMO LEER ESTE EXCEL",""),
 ("",""),
 ("Coincidencia CV original (%)","Cobertura ponderada de los requisitos de la oferta con el CV que ya tienes."),
 ("","Cada requisito lleva un peso según su importancia en la oferta. Un requisito vale 1,0 si"),
 ("","está demostrado en un bullet o en el resumen; 0,5 si sólo aparece listado en 'Competencias"),
 ("","técnicas' (señal débil para un reclutador y para un ATS); 0,0 si no lo tienes."),
 ("Coincidencia CV adaptado (%)","Lo mismo con el CV generado para esa oferta."),
 ("","La mejora viene SÓLO de sacar a un bullet o al resumen algo que ya sabías hacer."),
 ("","Ningún requisito que no tengas sube de 0,0: no se inventa experiencia."),
 ("Mejora (pp)","Diferencia en puntos porcentuales entre el CV adaptado y el original."),
 ("Mejora (%)","La misma diferencia en porcentaje relativo sobre la puntuación original."),
 ("Huecos principales","Requisitos con peso alto que NO cubres. Son los que te preguntarán en la entrevista."),
 ("",""),
 ("SALARIOS",""),
 ("Verde","Salario publicado en la propia oferta."),
 ("Amarillo","Estimación. La columna 'Base de la estimación salarial' explica de dónde sale cada una."),
 ("Fuentes de referencia","Guía Salarial Manfred 2026 (salarios en tecnología, España),"),
 ("","Informe de salarios en IA en España 2026 (Universidad VIU), Levels.fyi por empresa."),
 ("Filtro aplicado","Se han descartado las ofertas cuyo salario medio estimado no llega a 35.000 €."),
 ("",""),
 ("COVER LETTER Y CORREO A RRHH",""),
 ("Dónde están","Ya no se generan por adelantado. En el dashboard, dentro de cada oferta, hay un botón"),
 ("","para escribir la cover letter y otro para el correo a RRHH en el momento, sólo para las"),
 ("","ofertas que te interesen. El texto se guarda en el propio dashboard y se puede copiar,"),
 ("","descargar como .txt o regenerar."),
 ("",""),
 ("ÁMBITO",""),
 ("España","Contrato y empresa en España, en remoto."),
 ("Navarra / Gipuzkoa","Presencial o híbrido, dentro de las provincias que aceptas."),
 ("Internacional","Empresa de fuera de España que contrata en remoto y admite residir aquí."),
 ("",""),
 ("FILTROS DE BÚSQUEDA",""),
 ("Modalidad","100% remoto, salvo en Navarra y Gipuzkoa, donde se acepta híbrido y presencial."),
 ("","En las internacionales se ha verificado que la restricción geográfica permite residir en España."),
 ("Antigüedad","Ofertas publicadas en los 7 días anteriores al 2 de septiembre de 2026."),
 ("Idioma","El CV y la cover letter se generan en el idioma de la oferta."),
]
ws2.column_dimensions['A'].width = 30; ws2.column_dimensions['B'].width = 100
for i,(a,b) in enumerate(lineas, start=1):
    ws2.cell(i,1,a).font = Font(name="Arial", size=10, bold=(b=="" and a!=""))
    ws2.cell(i,2,b).font = Font(name="Arial", size=10)
    ws2.cell(i,2).alignment = Alignment(wrap_text=True, vertical="top")

wb.save(os.environ.get("RADAR_XLSX", "out/ofertas.xlsx"))
print("ok")
